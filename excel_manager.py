import os
import re
import shutil
import datetime
from typing import Dict, List, Any, Optional
import openpyxl
from openpyxl.utils import get_column_letter

# Dynamic portable path resolution across any user account and OS
USER_DOWNLOADS = os.path.join(os.path.expanduser("~"), "Downloads")
PRIMARY_FILE_PATH = os.path.join(USER_DOWNLOADS, "Monitoramento_Dores - FINAL.xlsx")
PROJECT_ROOT_PATH = os.path.join(os.path.dirname(__file__), "Monitoramento_Dores - FINAL.xlsx")
FALLBACK_FILE_PATH = os.path.join(os.path.dirname(__file__), "data", "Monitoramento_Dores - FINAL.xlsx")
BACKUP_DIR = os.path.join(os.path.dirname(__file__), "backups")
UPLOADS_DIR = os.path.join(os.path.dirname(__file__), "uploads")

COLUMN_MAPPING = {
    1: "id",
    2: "data_ocorrencia",
    3: "data_registro",
    4: "responsavel",
    5: "setor_responsavel",
    6: "setor_impactado",
    7: "filial",
    8: "categoria_dor",
    9: "nota_fiscal",
    10: "serie",
    11: "descricao_problema",
    12: "causa_raiz",
    13: "impacto_negocio",
    14: "qtd_nota_kg",
    15: "valor_notas",
    16: "prioridade",
    17: "status",
    18: "plano_acao",
    19: "prazo",
    20: "data_conclusao",
    21: "dias_aberto",
    22: "recorrente",
    23: "observacoes",
    24: "sla_dias",
    25: "data_limite_sla",
    26: "situacao_sla",
    27: "dias_para_vencimento",
    28: "responsavel_solucao",
    29: "numero_chamado",
    30: "ultima_atualizacao",
    31: "anexos",
}

FIELD_TO_COL = {v: k for k, v in COLUMN_MAPPING.items()}

PRIORITY_SLA_MAP = {
    "Crítica": 2,
    "Critica": 2,
    "Alta": 5,
    "Média": 10,
    "Media": 10,
    "Baixa": 15
}


def sanitize_filename(name: str) -> str:
    """Cleans filename to prevent path traversal and unsafe characters."""
    name = os.path.basename(name).strip()
    name = re.sub(r'[\\/*?:"<>|]', "", name)
    name = name.replace(" ", "_")
    return name or "anexo"


def format_file_size(size_bytes: int) -> str:
    """Formats raw bytes into human readable KB / MB."""
    if size_bytes < 1024:
        return f"{size_bytes} B"
    elif size_bytes < 1024 * 1024:
        return f"{round(size_bytes / 1024, 1)} KB"
    else:
        return f"{round(size_bytes / (1024 * 1024), 2)} MB"


def get_record_attachments(record_id: int) -> List[Dict[str, Any]]:
    """Lists all files attached to a specific record ID with rich metadata."""
    rec_dir = os.path.join(UPLOADS_DIR, str(record_id))
    if not os.path.exists(rec_dir):
        return []

    attachments = []
    for fname in sorted(os.listdir(rec_dir)):
        fpath = os.path.join(rec_dir, fname)
        if os.path.isfile(fpath):
            try:
                sz = os.path.getsize(fpath)
                mtime = datetime.datetime.fromtimestamp(os.path.getmtime(fpath)).strftime("%d/%m/%Y %H:%M")
                _, ext = os.path.splitext(fname)
                ext_lower = ext.lower()

                is_img = ext_lower in [".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".svg"]
                is_pdf = ext_lower == ".pdf"
                is_excel = ext_lower in [".xlsx", ".xls", ".csv"]
                is_word = ext_lower in [".doc", ".docx"]

                attachments.append({
                    "filename": fname,
                    "size_bytes": sz,
                    "size_formatted": format_file_size(sz),
                    "extension": ext_lower,
                    "uploaded_at": mtime,
                    "url": f"/api/attachments/{record_id}/{fname}",
                    "is_image": is_img,
                    "is_pdf": is_pdf,
                    "is_excel": is_excel,
                    "is_word": is_word
                })
            except Exception:
                pass
    return attachments


def sync_excel_attachments_column(record_id: int):
    """Syncs the attachment filenames into Column 31 (Anexos) of the Excel file."""
    try:
        file_path = get_active_excel_path()
        attachments = get_record_attachments(record_id)
        names = ", ".join([a["filename"] for a in attachments])

        wb = openpyxl.load_workbook(file_path)
        ws = wb["Monitoramento de Dores"]
        if ws.cell(1, 31).value != "Anexos":
            ws.cell(1, 31, "Anexos")

        for r in range(2, ws.max_row + 1):
            if r - 1 == record_id:
                ws.cell(r, 31, names if names else None)
                break

        wb.save(file_path)
        wb.close()
    except Exception as e:
        print(f"Notice: sync_excel_attachments_column: {e}")


def save_record_attachments(record_id: int, files: List[Any]) -> List[Dict[str, Any]]:
    """Saves a list of uploaded FileStorage objects to uploads/<record_id>/."""
    rec_dir = os.path.join(UPLOADS_DIR, str(record_id))
    os.makedirs(rec_dir, exist_ok=True)

    saved_names = []
    for file_obj in files:
        if not file_obj or not getattr(file_obj, "filename", None):
            continue
        original_name = file_obj.filename
        safe_name = sanitize_filename(original_name)
        base, ext = os.path.splitext(safe_name)

        target_path = os.path.join(rec_dir, safe_name)
        counter = 1
        while os.path.exists(target_path):
            safe_name = f"{base}_{counter}{ext}"
            target_path = os.path.join(rec_dir, safe_name)
            counter += 1

        file_obj.save(target_path)
        saved_names.append(safe_name)

    sync_excel_attachments_column(record_id)
    return get_record_attachments(record_id)


def delete_record_attachment(record_id: int, filename: str) -> bool:
    """Deletes an attachment file for a record and syncs with Excel."""
    rec_dir = os.path.join(UPLOADS_DIR, str(record_id))
    safe_name = os.path.basename(filename)
    target_path = os.path.join(rec_dir, safe_name)
    if os.path.exists(target_path) and os.path.isfile(target_path):
        os.remove(target_path)
        sync_excel_attachments_column(record_id)
        return True
    return False


def get_active_excel_path() -> str:
    """
    Returns the most appropriate Excel file path available on the computer:
    1. Downloads folder of current logged-in user
    2. Project root folder (if placed next to app.py)
    3. Project data/ folder (bundled working copy)
    """
    if os.path.exists(PRIMARY_FILE_PATH):
        return PRIMARY_FILE_PATH
    if os.path.exists(PROJECT_ROOT_PATH):
        return PROJECT_ROOT_PATH
    if os.path.exists(FALLBACK_FILE_PATH):
        return FALLBACK_FILE_PATH
    return FALLBACK_FILE_PATH


def create_backup(file_path: str) -> Optional[str]:
    """Creates a timestamped backup before modifying the Excel file."""
    try:
        os.makedirs(BACKUP_DIR, exist_ok=True)
        ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_name = f"backup_dores_{ts}.xlsx"
        backup_path = os.path.join(BACKUP_DIR, backup_name)
        shutil.copy2(file_path, backup_path)
        return backup_path
    except Exception as e:
        print(f"Warning: Failed to create backup: {e}")
        return None


def get_parameters() -> Dict[str, Any]:
    """Reads dropdown options and SLA parameters from Sheet 3 (Parâmetros)."""
    file_path = get_active_excel_path()
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.worksheets[2]  # Parâmetros sheet

    statuses = []
    prioridades = []
    recorrentes = []
    categorias = []
    setores = []
    filiais = []
    sla_map = {}

    for r in range(2, ws.max_row + 1):
        v1 = ws.cell(r, 1).value
        if v1 and str(v1).strip():
            statuses.append(str(v1).strip())

        v3 = ws.cell(r, 3).value
        if v3 and str(v3).strip():
            prioridades.append(str(v3).strip())

        v5 = ws.cell(r, 5).value
        if v5 and str(v5).strip():
            recorrentes.append(str(v5).strip())

        v7 = ws.cell(r, 7).value
        if v7 and str(v7).strip():
            categorias.append(str(v7).strip())

        v10 = ws.cell(r, 10).value
        v11 = ws.cell(r, 11).value
        if v10 and v11 is not None and str(v10).strip() in ["Baixa", "Média", "Media", "Alta", "Crítica", "Critica"]:
            try:
                sla_map[str(v10).strip()] = int(v11)
            except (ValueError, TypeError):
                pass

        v13 = ws.cell(r, 13).value
        if v13 and str(v13).strip():
            setores.append(str(v13).strip())

        v15 = ws.cell(r, 15).value
        if v15 and str(v15).strip():
            filiais.append(str(v15).strip())

    wb.close()

    if not statuses:
        statuses = ["Aberto", "Em Análise", "Em Tratativa", "Aguardando Terceiros", "Concluído", "Cancelado"]
    if not prioridades:
        prioridades = ["Baixa", "Média", "Alta", "Crítica"]
    if not recorrentes:
        recorrentes = ["Sim", "Não"]
    if not categorias:
        categorias = ["Fiscal", "Estoque", "Financeiro", "Faturamento", "Logística", "Comercial", "TI/Sistemas", "Qualidade", "Compliance", "Outros"]
    if not setores:
        setores = ["Fiscal", "Estoque", "Financeiro", "Faturamento", "Logística", "Comercial", "TI", "Suprimentos", "Qualidade", "Operações", "Outros"]
    if not filiais:
        filiais = ["Matupá", "Rondonópolis", "Paranaguá"]
    if not sla_map:
        sla_map = PRIORITY_SLA_MAP

    return {
        "status": list(dict.fromkeys(statuses)),
        "prioridade": list(dict.fromkeys(prioridades)),
        "recorrente": list(dict.fromkeys(recorrentes)),
        "categoria": list(dict.fromkeys(categorias)),
        "setor": list(dict.fromkeys(setores)),
        "filial": list(dict.fromkeys(filiais)),
        "sla_map": sla_map
    }


def parse_date(val: Any) -> Optional[datetime.date]:
    """Helper to safely parse any date/datetime/string to a datetime.date."""
    if not val:
        return None
    if isinstance(val, datetime.datetime):
        return val.date()
    if isinstance(val, datetime.date):
        return val
    if isinstance(val, str):
        val = val.strip()
        if not val or val.startswith("="):
            return None
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%Y"):
            try:
                return datetime.datetime.strptime(val, fmt).date()
            except ValueError:
                pass
    return None


def format_date_iso(val: Any) -> Optional[str]:
    d = parse_date(val)
    return d.isoformat() if d else None


def format_date_br(val: Any) -> Optional[str]:
    d = parse_date(val)
    return d.strftime("%d/%m/%Y") if d else None


def compute_derived_fields(record: Dict[str, Any], sla_map: Dict[str, int]) -> Dict[str, Any]:
    """Computes SLA days, Data Limite SLA, Situação SLA, Dias em Aberto, Dias para Vencimento."""
    today = datetime.date.today()
    reg_date = parse_date(record.get("data_registro"))
    conc_date = parse_date(record.get("data_conclusao"))
    prio = record.get("prioridade", "").strip() if record.get("prioridade") else ""
    status = record.get("status", "").strip() if record.get("status") else "Aberto"

    sla_days = sla_map.get(prio, PRIORITY_SLA_MAP.get(prio, 5))
    record["sla_dias"] = sla_days

    limite_date = None
    if reg_date and sla_days is not None:
        limite_date = reg_date + datetime.timedelta(days=sla_days)
    record["data_limite_sla"] = limite_date.isoformat() if limite_date else None
    record["data_limite_sla_br"] = limite_date.strftime("%d/%m/%Y") if limite_date else "-"

    if reg_date:
        if conc_date:
            dias_aberto = (conc_date - reg_date).days
        else:
            dias_aberto = (today - reg_date).days
        record["dias_aberto"] = max(0, dias_aberto)
    else:
        record["dias_aberto"] = 0

    if status == "Concluído":
        record["situacao_sla"] = "Concluído"
        record["dias_para_vencimento"] = None
    elif limite_date:
        diff_days = (limite_date - today).days
        record["dias_para_vencimento"] = diff_days
        if limite_date < today:
            record["situacao_sla"] = "Vencido"
        elif diff_days <= 2:
            record["situacao_sla"] = "A vencer"
        else:
            record["situacao_sla"] = "No prazo"
    else:
        record["situacao_sla"] = "Sem SLA"
        record["dias_para_vencimento"] = None

    return record


def get_all_records() -> List[Dict[str, Any]]:
    """Reads all valid rows from 'Monitoramento de Dores' sheet."""
    file_path = get_active_excel_path()
    params = get_parameters()
    sla_map = params["sla_map"]

    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb["Monitoramento de Dores"]

    records = []
    for r in range(2, ws.max_row + 1):
        c_reg = ws.cell(r, 3).value
        c_resp = ws.cell(r, 4).value
        c_desc = ws.cell(r, 11).value

        if c_reg is None and c_resp is None and c_desc is None:
            continue

        item = {"_excel_row": r}
        for col_idx, field_name in COLUMN_MAPPING.items():
            cell_val = ws.cell(r, col_idx).value
            if isinstance(cell_val, str) and cell_val.startswith("#"):
                cell_val = None
            item[field_name] = cell_val

        try:
            item["id"] = int(item["id"]) if item["id"] is not None else (r - 1)
        except (ValueError, TypeError):
            item["id"] = r - 1

        for df in ["data_ocorrencia", "data_registro", "prazo", "data_conclusao", "ultima_atualizacao"]:
            item[f"{df}_iso"] = format_date_iso(item.get(df))
            item[f"{df}_br"] = format_date_br(item.get(df)) or "-"

        for nf in ["nota_fiscal", "serie", "qtd_nota_kg", "numero_chamado"]:
            try:
                item[nf] = int(item[nf]) if item[nf] is not None and str(item[nf]).strip() != "" else None
            except (ValueError, TypeError):
                pass

        try:
            item["valor_notas"] = float(item["valor_notas"]) if item["valor_notas"] is not None else 0.0
        except (ValueError, TypeError):
            item["valor_notas"] = 0.0

        compute_derived_fields(item, sla_map)

        # Attachments metadata
        rec_attachments = get_record_attachments(item["id"])
        item["anexos"] = rec_attachments
        item["total_anexos"] = len(rec_attachments)
        item["anexos_nomes"] = ", ".join([a["filename"] for a in rec_attachments])

        records.append(item)

    wb.close()
    return records


def calculate_kpis(records: List[Dict[str, Any]]) -> Dict[str, Any]:
    """Calculates all executive metrics and KPIs, plus timeline and Pareto distributions."""
    total = len(records)
    abertas = 0
    concluidas = 0
    sla_vencido = 0
    sla_a_vencer = 0
    sla_no_prazo = 0
    valor_total = 0.0
    soma_dias_aberto = 0
    soma_dias_resolucao = 0
    count_concluidas_com_dias = 0
    count_abertas = 0
    total_com_anexos = 0
    total_arquivos = 0

    by_status = {}
    by_situacao_sla = {"No prazo": 0, "A vencer": 0, "Vencido": 0, "Concluído": 0}
    by_categoria = {}
    by_setor_responsavel = {}
    by_setor_impactado = {}
    by_filial = {}
    by_prioridade = {"Crítica": 0, "Alta": 0, "Média": 0, "Baixa": 0}
    by_recorrente = {"Sim": 0, "Não": 0}

    # Timeline dictionary by Month (YYYY-MM)
    timeline_dict = {}

    for rec in records:
        st = rec.get("status") or "Aberto"
        by_status[st] = by_status.get(st, 0) + 1

        if st == "Concluído":
            concluidas += 1
            if rec.get("dias_aberto") is not None:
                soma_dias_resolucao += rec["dias_aberto"]
                count_concluidas_com_dias += 1
        elif st != "Cancelado":
            abertas += 1
            if rec.get("dias_aberto") is not None:
                soma_dias_aberto += rec["dias_aberto"]
                count_abertas += 1

        sit_sla = rec.get("situacao_sla")
        if sit_sla in by_situacao_sla:
            by_situacao_sla[sit_sla] += 1

        if sit_sla == "Vencido":
            sla_vencido += 1
        elif sit_sla == "A vencer":
            sla_a_vencer += 1
        elif sit_sla == "No prazo":
            sla_no_prazo += 1

        val = rec.get("valor_notas") or 0.0
        valor_total += val

        cat = rec.get("categoria_dor") or "Não Informado"
        by_categoria[cat] = by_categoria.get(cat, 0) + 1

        s_resp = rec.get("setor_responsavel") or "Não Informado"
        by_setor_responsavel[s_resp] = by_setor_responsavel.get(s_resp, 0) + 1

        s_imp = rec.get("setor_impactado") or "Não Informado"
        by_setor_impactado[s_imp] = by_setor_impactado.get(s_imp, 0) + 1

        filial = rec.get("filial") or "Não Informada"
        by_filial[filial] = by_filial.get(filial, 0) + 1

        prio = rec.get("prioridade")
        if prio in by_prioridade:
            by_prioridade[prio] += 1
        elif prio:
            by_prioridade[prio] = by_prioridade.get(prio, 0) + 1

        rec_val = rec.get("recorrente")
        if rec_val in by_recorrente:
            by_recorrente[rec_val] += 1

        # Attachments count
        if rec.get("total_anexos", 0) > 0:
            total_com_anexos += 1
            total_arquivos += rec.get("total_anexos", 0)

        # Timeline calculation
        reg_iso = rec.get("data_registro_iso")
        if reg_iso:
            month_key = reg_iso[:7]  # YYYY-MM
            if month_key not in timeline_dict:
                timeline_dict[month_key] = {"registradas": 0, "concluidas": 0}
            timeline_dict[month_key]["registradas"] += 1
            if st == "Concluído":
                timeline_dict[month_key]["concluidas"] += 1

    tempo_medio_aberto = round(soma_dias_aberto / count_abertas, 1) if count_abertas > 0 else 0
    tempo_medio_resolucao = round(soma_dias_resolucao / count_concluidas_com_dias, 1) if count_concluidas_com_dias > 0 else 0
    taxa_resolucao = round((concluidas / total) * 100, 1) if total > 0 else 0

    # Sort timeline
    sorted_months = sorted(timeline_dict.keys())
    timeline_labels = [f"{m.split('-')[1]}/{m.split('-')[0]}" for m in sorted_months]
    timeline_registradas = [timeline_dict[m]["registradas"] for m in sorted_months]
    timeline_concluidas = [timeline_dict[m]["concluidas"] for m in sorted_months]

    # Top financial impacts (Pareto)
    sorted_financial = sorted(records, key=lambda x: x.get("valor_notas", 0.0), reverse=True)[:5]
    top_financeiro = [{
        "id": r["id"],
        "filial": r.get("filial", "-"),
        "categoria": r.get("categoria_dor", "-"),
        "descricao": r.get("descricao_problema", "-")[:45] + ("..." if len(r.get("descricao_problema", "")) > 45 else ""),
        "valor": r.get("valor_notas", 0.0),
        "status": r.get("status", "Aberto"),
        "prioridade": r.get("prioridade", "Média")
    } for r in sorted_financial if r.get("valor_notas", 0.0) > 0]

    return {
        "total_ocorrencias": total,
        "ocorrencias_abertas": abertas,
        "ocorrencias_concluidas": concluidas,
        "taxa_resolucao": taxa_resolucao,
        "sla_vencido": sla_vencido,
        "sla_a_vencer": sla_a_vencer,
        "sla_no_prazo": sla_no_prazo,
        "valor_total_notas": valor_total,
        "tempo_medio_aberto": tempo_medio_aberto,
        "tempo_medio_resolucao": tempo_medio_resolucao,
        "total_com_anexos": total_com_anexos,
        "total_arquivos": total_arquivos,
        "by_status": by_status,
        "by_situacao_sla": by_situacao_sla,
        "by_categoria": by_categoria,
        "by_setor_responsavel": by_setor_responsavel,
        "by_setor_impactado": by_setor_impactado,
        "by_filial": by_filial,
        "by_prioridade": by_prioridade,
        "by_recorrente": by_recorrente,
        "timeline": {
            "labels": timeline_labels,
            "registradas": timeline_registradas,
            "concluidas": timeline_concluidas
        },
        "top_financeiro": top_financeiro
    }


def add_record(form_data: Dict[str, Any], files: Optional[List[Any]] = None) -> Dict[str, Any]:
    """
    Inserts a new record into 'Monitoramento de Dores' sheet in the Excel file.
    Preserves existing formulas, sets values, writes formulas for calculated fields,
    and saves any attached files into uploads/<id>/.
    """
    file_path = get_active_excel_path()
    create_backup(file_path)

    wb = openpyxl.load_workbook(file_path)
    ws = wb["Monitoramento de Dores"]

    # Ensure Col 31 header
    if ws.cell(1, 31).value != "Anexos":
        ws.cell(1, 31, "Anexos")

    target_row = None
    for r in range(2, max(ws.max_row + 2, 1002)):
        c_reg = ws.cell(r, 3).value
        c_resp = ws.cell(r, 4).value
        c_desc = ws.cell(r, 11).value
        if c_reg is None and c_resp is None and c_desc is None:
            target_row = r
            break

    if target_row is None:
        target_row = ws.max_row + 1

    r = target_row
    new_id = r - 1

    dt_ocorrencia = parse_date(form_data.get("data_ocorrencia")) or datetime.date.today()
    dt_registro = parse_date(form_data.get("data_registro")) or datetime.date.today()
    dt_prazo = parse_date(form_data.get("prazo"))
    dt_conclusao = parse_date(form_data.get("data_conclusao"))
    dt_atualizacao = datetime.date.today()

    responsavel = (form_data.get("responsavel") or "").strip().upper()
    setor_responsavel = (form_data.get("setor_responsavel") or "").strip()
    setor_impactado = (form_data.get("setor_impactado") or "").strip()
    filial = (form_data.get("filial") or "").strip()
    categoria_dor = (form_data.get("categoria_dor") or "").strip()

    try:
        nota_fiscal = int(form_data.get("nota_fiscal")) if form_data.get("nota_fiscal") not in (None, "") else None
    except (ValueError, TypeError):
        nota_fiscal = None

    try:
        serie = int(form_data.get("serie")) if form_data.get("serie") not in (None, "") else None
    except (ValueError, TypeError):
        serie = None

    try:
        qtd_nota_kg = float(form_data.get("qtd_nota_kg")) if form_data.get("qtd_nota_kg") not in (None, "") else None
    except (ValueError, TypeError):
        qtd_nota_kg = None

    try:
        valor_notas = float(form_data.get("valor_notas")) if form_data.get("valor_notas") not in (None, "") else 0.0
    except (ValueError, TypeError):
        valor_notas = 0.0

    try:
        numero_chamado = int(form_data.get("numero_chamado")) if form_data.get("numero_chamado") not in (None, "") else None
    except (ValueError, TypeError):
        numero_chamado = form_data.get("numero_chamado")

    prioridade = form_data.get("prioridade") or "Média"
    status = form_data.get("status") or "Aberto"
    descricao = form_data.get("descricao_problema") or ""
    causa_raiz = form_data.get("causa_raiz") or ""
    impacto_negocio = form_data.get("impacto_negocio") or ""
    plano_acao = form_data.get("plano_acao") or ""
    recorrente = form_data.get("recorrente") or "Não"
    observacoes = form_data.get("observacoes") or ""
    resp_solucao = (form_data.get("responsavel_solucao") or "").strip()

    # Save uploaded files if provided
    saved_filenames = []
    if files:
        rec_dir = os.path.join(UPLOADS_DIR, str(new_id))
        os.makedirs(rec_dir, exist_ok=True)
        for f in files:
            if not f or not getattr(f, "filename", None):
                continue
            s_name = sanitize_filename(f.filename)
            base, ext = os.path.splitext(s_name)
            t_path = os.path.join(rec_dir, s_name)
            cnt = 1
            while os.path.exists(t_path):
                s_name = f"{base}_{cnt}{ext}"
                t_path = os.path.join(rec_dir, s_name)
                cnt += 1
            f.save(t_path)
            saved_filenames.append(s_name)

    # Set cell values
    ws.cell(r, 1, f'=IF(C{r}="","",ROW()-1)')
    ws.cell(r, 2, datetime.datetime.combine(dt_ocorrencia, datetime.time.min))
    ws.cell(r, 3, datetime.datetime.combine(dt_registro, datetime.time.min))
    ws.cell(r, 4, responsavel)
    ws.cell(r, 5, setor_responsavel)
    ws.cell(r, 6, setor_impactado)
    ws.cell(r, 7, filial)
    ws.cell(r, 8, categoria_dor)
    ws.cell(r, 9, nota_fiscal)
    ws.cell(r, 10, serie)
    ws.cell(r, 11, descricao)
    ws.cell(r, 12, causa_raiz)
    ws.cell(r, 13, impacto_negocio)
    ws.cell(r, 14, qtd_nota_kg)
    ws.cell(r, 15, valor_notas)
    ws.cell(r, 16, prioridade)
    ws.cell(r, 17, status)
    ws.cell(r, 18, plano_acao)
    ws.cell(r, 19, datetime.datetime.combine(dt_prazo, datetime.time.min) if dt_prazo else None)
    ws.cell(r, 20, datetime.datetime.combine(dt_conclusao, datetime.time.min) if dt_conclusao else None)
    ws.cell(r, 21, f'=IF(C{r}="","",IF(T{r}="",TODAY()-C{r},T{r}-C{r}))')
    ws.cell(r, 22, recorrente)
    ws.cell(r, 23, observacoes)
    ws.cell(r, 24, f'=IF(P{r}="","",IFERROR(VLOOKUP(P{r},Parâmetros!$J$2:$K$5,2,FALSE),""))')
    ws.cell(r, 25, f'=IF(OR(C{r}="",X{r}=""),"",C{r}+X{r})')
    ws.cell(r, 26, f'=IF(C{r}="","",IF(Q{r}="Concluído","Concluído",IF(Y{r}<TODAY(),"Vencido",IF(Y{r}-TODAY()<=2,"A vencer","No prazo"))))')
    ws.cell(r, 27, f'=IF(OR(Y{r}="",Q{r}="Concluído"),"",Y{r}-TODAY())')
    ws.cell(r, 28, resp_solucao)
    ws.cell(r, 29, numero_chamado)
    ws.cell(r, 30, datetime.datetime.combine(dt_atualizacao, datetime.time.min))
    ws.cell(r, 31, ", ".join(saved_filenames) if saved_filenames else None)

    try:
        wb.save(file_path)
    except PermissionError:
        wb.close()
        return {
            "success": False,
            "message": "A planilha está aberta no Excel! Por favor, feche o Excel para salvar a nova pendência."
        }
    except Exception as e:
        wb.close()
        return {
            "success": False,
            "message": f"Erro ao salvar na planilha Excel: {str(e)}"
        }

    wb.close()

    if file_path == PRIMARY_FILE_PATH and os.path.exists(os.path.dirname(FALLBACK_FILE_PATH)):
        try:
            shutil.copy2(PRIMARY_FILE_PATH, FALLBACK_FILE_PATH)
        except Exception:
            pass

    return {
        "success": True,
        "row": r,
        "id": new_id,
        "attachments_count": len(saved_filenames),
        "message": f"Pendência #{new_id} cadastrada com sucesso diretamente no arquivo Excel!"
    }


def update_record(record_id: int, form_data: Dict[str, Any]) -> Dict[str, Any]:
    """Updates an existing record in the Excel file by its ID."""
    file_path = get_active_excel_path()
    create_backup(file_path)

    wb = openpyxl.load_workbook(file_path)
    ws = wb["Monitoramento de Dores"]

    target_row = None
    for r in range(2, ws.max_row + 1):
        if r - 1 == record_id:
            target_row = r
            break

    if not target_row:
        wb.close()
        return {"success": False, "message": f"Registro ID {record_id} não encontrado."}

    r = target_row

    if "responsavel" in form_data and form_data["responsavel"] is not None:
        ws.cell(r, 4, str(form_data["responsavel"]).strip().upper())

    if "setor_responsavel" in form_data and form_data["setor_responsavel"] is not None:
        ws.cell(r, 5, str(form_data["setor_responsavel"]).strip())

    if "setor_impactado" in form_data and form_data["setor_impactado"] is not None:
        ws.cell(r, 6, str(form_data["setor_impactado"]).strip())

    if "filial" in form_data and form_data["filial"] is not None:
        ws.cell(r, 7, str(form_data["filial"]).strip())

    if "categoria_dor" in form_data and form_data["categoria_dor"] is not None:
        ws.cell(r, 8, str(form_data["categoria_dor"]).strip())

    if "descricao_problema" in form_data and form_data["descricao_problema"] is not None:
        ws.cell(r, 11, str(form_data["descricao_problema"]))

    if "causa_raiz" in form_data and form_data["causa_raiz"] is not None:
        ws.cell(r, 12, str(form_data["causa_raiz"]))

    if "impacto_negocio" in form_data and form_data["impacto_negocio"] is not None:
        ws.cell(r, 13, str(form_data["impacto_negocio"]))

    if "prioridade" in form_data and form_data["prioridade"]:
        ws.cell(r, 16, str(form_data["prioridade"]).strip())

    if "status" in form_data and form_data["status"]:
        new_status = str(form_data["status"]).strip()
        ws.cell(r, 17, new_status)
        if new_status == "Concluído" and not ws.cell(r, 20).value and not form_data.get("data_conclusao"):
            ws.cell(r, 20, datetime.datetime.combine(datetime.date.today(), datetime.time.min))

    if "plano_acao" in form_data:
        ws.cell(r, 18, form_data["plano_acao"] or "")

    if "prazo" in form_data:
        dt_prazo = parse_date(form_data["prazo"])
        ws.cell(r, 19, datetime.datetime.combine(dt_prazo, datetime.time.min) if dt_prazo else None)

    if "data_conclusao" in form_data:
        dt_conc = parse_date(form_data["data_conclusao"])
        ws.cell(r, 20, datetime.datetime.combine(dt_conc, datetime.time.min) if dt_conc else None)

    if "recorrente" in form_data and form_data["recorrente"]:
        ws.cell(r, 22, str(form_data["recorrente"]).strip())

    if "observacoes" in form_data:
        ws.cell(r, 23, form_data["observacoes"] or "")

    if "responsavel_solucao" in form_data:
        ws.cell(r, 28, (form_data["responsavel_solucao"] or "").strip())

    if "nota_fiscal" in form_data:
        try:
            ws.cell(r, 9, int(form_data["nota_fiscal"]) if form_data["nota_fiscal"] not in (None, "") else None)
        except (ValueError, TypeError):
            pass

    if "serie" in form_data:
        try:
            ws.cell(r, 10, int(form_data["serie"]) if form_data["serie"] not in (None, "") else None)
        except (ValueError, TypeError):
            pass

    if "qtd_nota_kg" in form_data:
        try:
            ws.cell(r, 14, float(form_data["qtd_nota_kg"]) if form_data["qtd_nota_kg"] not in (None, "") else None)
        except (ValueError, TypeError):
            pass

    if "valor_notas" in form_data:
        try:
            ws.cell(r, 15, float(form_data["valor_notas"]) if form_data["valor_notas"] not in (None, "") else 0.0)
        except (ValueError, TypeError):
            pass

    if "numero_chamado" in form_data:
        try:
            ws.cell(r, 29, int(form_data["numero_chamado"]) if form_data["numero_chamado"] not in (None, "") else None)
        except (ValueError, TypeError):
            ws.cell(r, 29, form_data["numero_chamado"])

    # Ensure Col 31 header
    if ws.cell(1, 31).value != "Anexos":
        ws.cell(1, 31, "Anexos")

    # Update attachment names in Col 31
    existing_attachments = get_record_attachments(record_id)
    if existing_attachments:
        ws.cell(r, 31, ", ".join([a["filename"] for a in existing_attachments]))

    # Formulas
    ws.cell(r, 1, f'=IF(C{r}="","",ROW()-1)')
    ws.cell(r, 21, f'=IF(C{r}="","",IF(T{r}="",TODAY()-C{r},T{r}-C{r}))')
    ws.cell(r, 24, f'=IF(P{r}="","",IFERROR(VLOOKUP(P{r},Parâmetros!$J$2:$K$5,2,FALSE),""))')
    ws.cell(r, 25, f'=IF(OR(C{r}="",X{r}=""),"",C{r}+X{r})')
    ws.cell(r, 26, f'=IF(C{r}="","",IF(Q{r}="Concluído","Concluído",IF(Y{r}<TODAY(),"Vencido",IF(Y{r}-TODAY()<=2,"A vencer","No prazo"))))')
    ws.cell(r, 27, f'=IF(OR(Y{r}="",Q{r}="Concluído"),"",Y{r}-TODAY())')
    ws.cell(r, 30, datetime.datetime.combine(datetime.date.today(), datetime.time.min))

    try:
        wb.save(file_path)
    except PermissionError:
        wb.close()
        return {
            "success": False,
            "message": "A planilha está aberta no Excel! Por favor, feche o Excel para permitir salvar as alterações."
        }
    except Exception as e:
        wb.close()
        return {"success": False, "message": f"Erro ao salvar planilha: {str(e)}"}

    wb.close()

    if file_path == PRIMARY_FILE_PATH and os.path.exists(os.path.dirname(FALLBACK_FILE_PATH)):
        try:
            shutil.copy2(PRIMARY_FILE_PATH, FALLBACK_FILE_PATH)
        except Exception:
            pass

    return {
        "success": True,
        "row": r,
        "id": record_id,
        "message": f"Registro #{record_id} atualizado com sucesso no Excel!"
    }


def quick_complete_record(record_id: int, obs: str = "") -> Dict[str, Any]:
    """Quick 1-click completion of a record."""
    return update_record(record_id, {
        "status": "Concluído",
        "data_conclusao": datetime.date.today().isoformat(),
        "observacoes": obs
    })
